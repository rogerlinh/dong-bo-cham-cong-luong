# -*- coding: utf-8 -*-
"""
Đồng bộ bảng chấm công của nhiều dự án (mỗi bên một biểu mẫu khác nhau) về 1 file Excel chuẩn theo
cấu trúc template: DM + 4 tab chấm công
    bang_cham_cong_thang_nha_may, bang_cham_cong_tuan_nha_may  -> doanh thu SK thu của nhà máy
    bang_cham_cong_thang_SK,      bang_cham_cong_tuan_SK       -> chi phí SK trả người lao động
+ 1 tab cho từng dự án (để đối chiếu với file gốc), cham_cong_chi_tiet (dạng dọc) và kiem_tra (cảnh báo).

Dự án có 2 bảng công (nhà máy và SK) thì nạp mỗi bảng vào tab tương ứng; dự án chỉ có 1 bảng công thì
bảng đó được nạp vào cả tab nhà máy và tab SK. Dự án chấm công theo tháng chỉ nạp vào 2 tab tháng, dự án
chấm công theo tuần (kỳ công không quá 7 ngày) chỉ nạp vào 2 tab tuần.

Chỉ đọc tab CHẤM CÔNG và tab BẢNG KÊ của file nguồn, KHÔNG đọc bảng lương: bảng lương trong file nguồn là
bản đã tổng hợp, bảng lương mới (bang_luong, doanh_thu của template) sẽ được tạo lại từ các tab chấm công.
Bảng kê (bảng tính tiền với nhà máy) chỉ dùng để bổ sung thông tin NV và đối chiếu với công nhà máy.

Cách dùng (dòng lệnh):
    python dong_bo_luong.py <file1.xlsx> <file2.xlsx> ... -o KET_QUA.xlsx
    python dong_bo_luong.py --thu-muc "C:\\...\\template luong" -o KET_QUA.xlsx

Mỗi biểu mẫu được nhận diện tự động theo tên sheet (xem ADAPTERS).
Muốn thêm dự án mới có biểu mẫu khác: viết thêm 1 hàm doc_mau_X() và khai báo trong ADAPTERS.
"""
from __future__ import annotations

import argparse
import datetime as dt
import glob
import os
import re
import sys
import warnings
import zipfile
from collections import Counter, defaultdict
from copy import copy
from dataclasses import dataclass, field

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

# --------------------------------------------------------------------------------------
# Danh mục loại giờ chuẩn (giữ nguyên tên như template; 2 loại cuối là bổ sung)
# --------------------------------------------------------------------------------------
LOAI_GIO = [
    # code, tên trong template (giữ nguyên ký tự xuống dòng / 2 dấu cách như template)
    ("X", "HC Ngày 100%"),
    ("D", "Tăng ca ngày\n150%"),
    ("OT", "HC  đêm 130%"),
    ("OD", "Tăng ca đêm 200%"),
    ("OCN", "Tăng ca ngày chủ nhật"),
    ("OL", "Tăng ca ngày lễ"),
    ("L", "Nghỉ lễ có lương"),
    ("P", "Nghỉ phép có lương"),
    # bổ sung (template chưa có) – để cuối để 8 dòng đầu DM đúng vị trí template
    ("OCND", "Tăng ca đêm chủ nhật"),
    ("OLD", "Tăng ca đêm ngày lễ"),
]
# thứ tự hiển thị các dòng giờ phụ trong bảng chấm công tháng/tuần
THU_TU_GIO_PHU = ["OCN", "OCND", "OL", "OLD", "L", "P"]
TEN_LOAI = dict(LOAI_GIO)
# template bang_cham_cong_tuan ghi tên loại giờ trên 1 dòng, 1 dấu cách (vd 'Tăng ca ngày 150%', 'HC đêm 130%')
TEN_LOAI_TUAN = {k: " ".join(v.split()) for k, v in TEN_LOAI.items()}
LOAI_CHINH = ["X", "D", "OT", "OD"]           # 4 dòng luôn có trong bảng chấm công tháng/tuần (như template)
LOAI_LAM = ["X", "D", "OT", "OD", "OCN", "OCND", "OL", "OLD"]   # các loại giờ làm việc (đối chiếu bảng kê)
LOAI_TANG_CA = ["D", "OD", "OCN", "OCND", "OL", "OLD"]

# Mức cảnh báo
LOI, NGHI_NGO, LUU_Y = "Lỗi", "Nghi ngờ", "Lưu ý"
THU_TU_MUC = {LOI: 0, NGHI_NGO: 1, LUU_Y: 2}

# Ngưỡng kiểm tra (chỉnh tại đây nếu chính sách khác)
NGUONG = {
    "gio_ngay_nghi_ngo": 16,      # tổng giờ 1 ngày > 16h -> nghi ngờ
    "gio_ngay_loi": 24,           # > 24h/ngày -> lỗi chắc chắn
    "tang_ca_thang": 40,          # BLLĐ 2019 Điều 107: làm thêm không quá 40 giờ/tháng
    "sai_so_gio": 0.01,           # lệch giờ chấm công / bảng kê > 0,01h coi là lệch
}
LOI_CONG_THUC = ("#REF!", "#N/A", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")


@dataclass
class NhanVien:
    ma_nv: str
    ten: str = ""
    ngay_vao: dt.date | None = None
    ngay_nghi: dt.date | None = None
    bo_phan: str = ""
    chuc_danh: str = ""
    ncc: str = ""
    ma_the: str = ""
    tinh: str = ""
    cccd: str = ""
    stk: str = ""


@dataclass
class DuAn:
    ma: str                      # "CTY A"
    file: str
    mau: str                     # tên biểu mẫu nhận diện
    ky: str = ""                 # "T08/2026"
    nguon_cc: str = ""           # sheet chấm công SK đã đọc
    tu_ngay: dt.date | None = None
    den_ngay: dt.date | None = None
    nv: dict = field(default_factory=dict)             # ma_nv -> NhanVien, theo thứ tự trong bảng công
    # công SK (tính chi phí trả người lao động): (ma, date) -> {code: h}
    cham_cong: dict = field(default_factory=lambda: defaultdict(lambda: defaultdict(float)))
    # công nhà máy (tính doanh thu); None = dự án chỉ có 1 bảng công, dùng chung với công SK
    cham_cong_nm: dict | None = None
    nguon_nm: str = ""           # sheet chấm công nhà máy đã đọc
    ky_hieu_ngay: dict = field(default_factory=dict)   # (ma, date) -> "N"/"TD"/...
    trung_ma: dict = field(default_factory=dict)       # ma_nv -> số khối/dòng NV trùng mã trong bảng công
    don_gia_dv: dict = field(default_factory=dict)     # code -> đơn giá dịch vụ (từ bảng kê, nếu có)
    nguon_don_gia: str = ""
    nguon_doi_chieu: str = ""    # sheet dùng để đối chiếu giờ (bảng kê / bảng công phía khách hàng)
    bang_ke: dict | None = None        # ma_nv -> {code: h}: bảng kê theo từng NV
    bang_ke_tong: dict | None = None   # code -> h: bảng kê chỉ có tổng theo loại giờ
    van_de: list = field(default_factory=list)         # [(muc, ma_nv, loai, noi_dung)]
    o_loi: list = field(default_factory=list)          # [(ma_nv, sheet, ô, giá trị lỗi)]
    loai_mau: list = field(default_factory=list)       # loại giờ bảng công gốc luôn có dòng riêng cho mỗi NV

    def bao(self, muc, ma_nv, loai, noi_dung):
        """Ghi 1 cảnh báo. muc: LOI (sai chắc chắn) / NGHI_NGO (cần kiểm tra) / LUU_Y (thông tin)."""
        self.van_de.append((muc, ma_nv, loai, noi_dung))

    @property
    def hai_bang(self):
        """True khi dự án có riêng bảng công nhà máy và bảng công SK."""
        return self.cham_cong_nm is not None

    @property
    def cc_nm(self):
        """Công nhà máy (dự án 1 bảng công: chính là công SK)."""
        return self.cham_cong if self.cham_cong_nm is None else self.cham_cong_nm

    @property
    def theo_tuan(self):
        """Dự án chấm công theo tuần (kỳ công trong bảng chấm công không quá 7 ngày) – chỉ nạp vào 2 tab tuần;
        còn lại là chấm công theo tháng – chỉ nạp vào 2 tab tháng."""
        return bool(self.tu_ngay and self.den_ngay and (self.den_ngay - self.tu_ngay).days < 7)

    def moi_bang_nm(self, sheet):
        self.cham_cong_nm = defaultdict(lambda: defaultdict(float))
        self.nguon_nm = sheet
        return self.cham_cong_nm

    def tong_gio_nv(self, cc=None):
        """ma_nv -> {code: tổng giờ trong kỳ} (chỉ NV trong danh sách). cc mặc định là công SK."""
        t = defaultdict(lambda: defaultdict(float))
        for (m, _), g in (self.cham_cong if cc is None else cc).items():
            if m in self.nv:
                for c, h in g.items():
                    t[m][c] += h
        return t


# --------------------------------------------------------------------------------------
# Hàm tiện ích
# --------------------------------------------------------------------------------------
def so(v) -> float:
    """Chuyển giá trị ô về số; text '0,5' / '12,35' / '1.000' được xử lý; còn lại -> 0."""
    if v is None or isinstance(v, bool):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("\xa0", "").replace(" ", "")
    s = re.sub(r"(?i)(vnđ|vnd|đồng|đ|giờ|h)$", "", s)  # đơn vị dính sau số: 2.500.000đ, 8h
    if not s:
        return 0.0
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})+(,\d+)?", s):     # 1.000.000 / 1.000,5 (kiểu Việt)
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(,\d{3})+(\.\d+)?", s):   # 1,000,000 / 1,000.5 (kiểu Anh)
        s = s.replace(",", "")
    else:
        s = s.replace(",", ".")                              # 7,5 giờ
    try:
        return float(s)
    except ValueError:
        return 0.0


def ma(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def ten_sach(v) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip()


def chuan(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def ci(col: str) -> int:
    return column_index_from_string(col) - 1


def o(row, col: str):
    i = ci(col)
    return row[i] if i < len(row) else None


def la_ngay(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str):
        try:
            m = re.match(r"\s*(\d{4})[/-](\d{1,2})[/-](\d{1,2})", v)
            if m:
                return dt.date(int(m[1]), int(m[2]), int(m[3]))
            m = re.match(r"\s*(\d{1,2})/(\d{1,2})/(\d{4})", v)
            if m:
                return dt.date(int(m[3]), int(m[2]), int(m[1]))
        except ValueError:          # ngày không tồn tại, vd 31/02
            return None
    return None


def co_chu_han(s: str) -> bool:
    return bool(re.search(r"[\u4e00-\u9fff]", s or ""))


def doc_sheet(wb, ten, min_row=1, max_row=None, dung_khi_trong=60):
    """Đọc các dòng (values) của sheet; tự dừng khi gặp chuỗi dài dòng trống (sheet không có dimension)."""
    ws = wb[ten]
    rows, trong = [], 0
    for r in ws.iter_rows(min_row=min_row, max_row=max_row, values_only=True):
        rows.append(r)
        if all(v in (None, "") for v in r):
            trong += 1
            if trong >= dung_khi_trong:
                break
        else:
            trong = 0
    while rows and all(v in (None, "") for v in rows[-1]):
        rows.pop()
    return rows


def ky_tu_ten_file(path):
    m = re.search(r"T(\d{1,2})[._-](\d{4})", os.path.basename(path))
    return (int(m[1]), int(m[2])) if m else None


def ma_du_an_tu_file(path, mac_dinh=None):
    m = re.search(r"CTY[_ ]?([A-Za-z0-9]+)", os.path.basename(path))
    if m:
        return f"CTY {m[1].upper()}"
    # không có 'CTY...' trong tên file: dùng chính tên file để 2 file khác nhau không bị gộp nhầm 1 mã
    return os.path.splitext(os.path.basename(path))[0][:40]


def them_nv(da: DuAn, m, **tt):
    """Thêm NV (nếu chưa có) và bổ sung các thông tin còn trống; không ghi đè thông tin đã có."""
    nv = da.nv.get(m)
    if nv is None:
        nv = da.nv[m] = NhanVien(ma_nv=m)
    for k, v in tt.items():
        if v not in (None, "") and getattr(nv, k) in (None, ""):
            setattr(nv, k, v)
    return nv


def nap_kieu_khoi(da: DuAn, rows, cot_ma, cot_nhan, cot_ngay: dict, map_nhan, sheet, dich=None):
    """Bảng công dạng khối: mỗi NV nhiều dòng (mỗi dòng 1 loại giờ), mã NV ở dòng đầu khối (hoặc mọi dòng).
    dich: bảng công nhận giờ (mặc định công SK da.cham_cong; bảng công nhà máy truyền da.cham_cong_nm).
    Trả về {mã: dòng đầu khối} để bộ đọc lấy thông tin NV."""
    la_sk = dich is None
    dich = da.cham_cong if dich is None else dich
    ma_hien_tai, dau_khoi, nhan_khoi = None, {}, set()
    for r in rows:
        m = ma(o(r, cot_ma))
        nhan = chuan(o(r, cot_nhan))
        # khối mới: mã khác mã đang đọc, hoặc cùng mã nhưng nhãn giờ lặp lại (2 khối liền nhau trùng mã)
        if m and (m != ma_hien_tai or (nhan and nhan in nhan_khoi)):
            if m in dau_khoi:                          # mã đã có ở 1 khối trước đó -> trùng mã
                if la_sk:
                    da.trung_ma[m] = da.trung_ma.get(m, 1) + 1
                else:
                    da.bao(LOI, m, "Trùng mã NV", f"Mã xuất hiện ở nhiều khối NV trong '{sheet}' – giờ đã bị cộng dồn vào 1 NV")
            else:
                dau_khoi[m] = r
            ma_hien_tai, nhan_khoi = m, set()
        if nhan:
            nhan_khoi.add(nhan)
        if not nhan or ma_hien_tai is None:
            continue
        code = map_nhan.get(nhan)
        if code is None:
            da.bao(NGHI_NGO, ma_hien_tai, "Nhãn giờ lạ", f"'{sheet}': không nhận ra loại giờ '{o(r, cot_nhan)}' - bỏ qua dòng này")
            continue
        for col, ngay in cot_ngay.items():
            v = o(r, col)
            if isinstance(v, str) and v.strip().upper() in LOI_CONG_THUC:
                da.o_loi.append((ma_hien_tai, sheet, f"{col} ({ngay:%d/%m})", v.strip()))
                continue
            if isinstance(v, str) and v.strip() and so(v) == 0 and not re.fullmatch(r"[\d.,\s]+", v):
                if la_sk or (ma_hien_tai, ngay) not in da.ky_hieu_ngay:
                    da.ky_hieu_ngay[(ma_hien_tai, ngay)] = v.strip()
                continue
            h = so(v)
            if h:
                dich[(ma_hien_tai, ngay)][code] += h
    return dau_khoi


def hang_co_nhieu_ngay(rows, so_dong=15, tu_cot=0):
    """Vị trí dòng tiêu đề ngày (dòng có nhiều ô ngày nhất) + {cột: ngày}."""
    i = max(range(min(so_dong, len(rows))), key=lambda k: sum(isinstance(v, dt.datetime) for v in rows[k][tu_cot:]))
    return i, {get_column_letter(j + 1): v.date() for j, v in enumerate(rows[i])
               if isinstance(v, dt.datetime) and j >= tu_cot}


def doc_bang_ke_theo_nv(wb, sheet, hang_dau, cot_ma, cot_gio: dict, cot_don_gia: dict | None = None,
                        dung_khi=None):
    """Bảng kê dạng 1 dòng/NV. cot_gio: {code: cột giờ, hoặc tuple các cột cộng dồn}; cot_don_gia: {code: cột đơn giá (lấy giá phổ biến)}.
    Trả về ({mã: {code: giờ}}, {mã: dòng}, {code: đơn giá})."""
    gio, dong, dem_gia = {}, {}, defaultdict(Counter)
    for r in doc_sheet(wb, sheet)[hang_dau - 1:]:
        if dung_khi and dung_khi(r):
            break
        m = ma(o(r, cot_ma))
        if not m:
            continue
        g = gio.setdefault(m, defaultdict(float))
        for code, col in cot_gio.items():
            g[code] += sum(so(o(r, c)) for c in (col if isinstance(col, tuple) else (col,)))
        dong.setdefault(m, r)
        for code, col in (cot_don_gia or {}).items():
            if so(o(r, col)):
                dem_gia[code][so(o(r, col))] += 1
    return gio, dong, {c: d.most_common(1)[0][0] for c, d in dem_gia.items()}


# --------------------------------------------------------------------------------------
# MẪU A: chấm công 'BCC LĐ THUÊ LẠI mm.yyyy' (4 dòng/NV); đối chiếu 'Bảng kê' (tổng giờ theo loại)
# --------------------------------------------------------------------------------------
def sheet_cung_ky(da: DuAn, sheet, ngay_ds, ma_ds, vai_tro):
    """Bảng công thứ 2 (vd bảng công SK) chỉ dùng khi cùng kỳ với bảng công chính và có NV chung.
    Không dùng được thì ghi lưu ý (kỳ của sheet, số NV chung) và trả về False."""
    ngay_ds = [n for n in ngay_ds if n]
    chung = len(set(ma_ds) & set(da.nv))
    trong_ky = [n for n in ngay_ds if da.tu_ngay and da.tu_ngay <= n <= da.den_ngay]
    if ngay_ds and chung and len(trong_ky) * 2 >= len(ngay_ds):
        return True
    ky = f"{min(ngay_ds):%d/%m/%Y} – {max(ngay_ds):%d/%m/%Y}" if ngay_ds else "không có ngày"
    da.bao(LUU_Y, "", "Bỏ qua bảng công khác kỳ",
           f"Sheet '{sheet}' ({vai_tro}) là dữ liệu kỳ {ky}, {len(set(ma_ds))} NV, chung {chung} NV với '{da.nguon_cc}' "
           f"(kỳ {da.tu_ngay:%d/%m/%Y} – {da.den_ngay:%d/%m/%Y}) – không dùng. "
           + ("Dự án coi như chỉ có 1 bảng công." if da.cham_cong_nm is None else ""))
    return False


def nhan_dien_A(sheets):
    return any(s.upper().startswith("BCC") for s in sheets)


def loai_tu_dien_giai(t):
    """Dòng diễn giải của bảng kê công nợ -> loại giờ ('Lương giờ làm ban đêm ngày thường' -> OT ...)."""
    t = chuan(t)
    dem = "đêm" in t
    if "lễ" in t:
        return "OLD" if dem else "OL"
    if "nghỉ" in t:
        return "OCND" if dem else "OCN"
    if "tăng ca" in t or "làm thêm" in t:
        return "OD" if dem else "D"
    if "ngày thường" in t or "giờ làm" in t:
        return "OT" if dem else "X"
    return None


def doc_mau_A(path, wb) -> DuAn:
    bcc = next(s for s in wb.sheetnames if s.upper().startswith("BCC"))
    da = DuAn(ma=ma_du_an_tu_file(path, "Dự án A"), file=path, mau="Mẫu A (1 bảng công: BCC; đối chiếu Bảng kê)", nguon_cc=bcc)
    rows = doc_sheet(wb, bcc)
    hang_ngay, cot_ngay = hang_co_nhieu_ngay(rows)
    map_nhan = {"hc ngày 100%": "X", "tăng ca ngày 150%": "D", "hc đêm 130%": "OT", "tăng ca đêm 200%": "OD",
                "tăng ca ngày nghỉ": "OCN", "tăng ca ngày lễ": "OL"}
    du_lieu = [r for r in rows[hang_ngay + 1:] if chuan(o(r, "A")) not in ("tổng", "tổng cộng")]
    for m, r in nap_kieu_khoi(da, du_lieu, "B", "I", cot_ngay, map_nhan, bcc).items():
        them_nv(da, m, ten=ten_sach(o(r, "C")), bo_phan=ten_sach(o(r, "D")), chuc_danh=ten_sach(o(r, "E")),
                ncc=ten_sach(o(r, "F")), ngay_vao=la_ngay(o(r, "G")), ngay_nghi=la_ngay(o(r, "H")))
    da.tu_ngay, da.den_ngay = min(cot_ngay.values()), max(cot_ngay.values())

    # Bảng kê công nợ: mỗi dòng 1 loại giờ (B/C diễn giải, E số lượng, F đơn giá)
    if "Bảng kê" in wb.sheetnames:
        rows = doc_sheet(wb, "Bảng kê")
        hang_td = next((i for i, r in enumerate(rows) if chuan(o(r, "A")) == "stt"), None)
        if hang_td is not None:
            da.bang_ke_tong = defaultdict(float)
            for r in rows[hang_td + 1:]:
                if chuan(o(r, "A")).startswith("tổng"):
                    break
                code = loai_tu_dien_giai(o(r, "B") or o(r, "C"))
                if not code:
                    continue
                da.bang_ke_tong[code] += so(o(r, "E"))
                if so(o(r, "F")) and (so(o(r, "E")) or code not in da.don_gia_dv):
                    da.don_gia_dv[code] = so(o(r, "F"))
            da.nguon_doi_chieu = da.nguon_don_gia = "Bảng kê"
    # 'DL chấm công': chi tiết chấm công từng ngày (dữ liệu máy chấm công). Bảng kê tính theo BCC nên BCC là
    # bảng công duy nhất; DL chấm công chỉ được báo nếu khác kỳ (thường là dữ liệu tháng cũ còn sót).
    for sh in wb.sheetnames:
        if chuan(sh).startswith("dl chấm công"):
            rows = doc_sheet(wb, sh, min_row=4)
            if any(o(r, "A") for r in rows):
                sheet_cung_ky(da, sh, [la_ngay(o(r, "E")) for r in rows], [ma(o(r, "A")) for r in rows if o(r, "A")],
                              "chi tiết chấm công")
    return da


# --------------------------------------------------------------------------------------
# MẪU B: chấm công 'Bảng công SK' (6 dòng/NV); đối chiếu 'Tính phí' (theo giờ ' Bảng công nhà máy')
# --------------------------------------------------------------------------------------
def nhan_dien_B(sheets):
    return "Bảng công SK" in sheets


def doc_mau_B(path, wb) -> DuAn:
    da = DuAn(ma=ma_du_an_tu_file(path, "Dự án B"), file=path, mau="Mẫu B (1 bảng công: Bảng công SK; đối chiếu Tính phí)",
              nguon_cc="Bảng công SK", loai_mau=["X", "OT", "D", "OD", "OCN", "OCND"])   # mỗi NV 6 dòng như file gốc
    rows = doc_sheet(wb, "Bảng công SK")
    hang_ngay, cot_ngay = hang_co_nhieu_ngay(rows, 10, ci("I"))
    map_nhan = {"ca ngày": "X", "ca đêm": "OT", "tăng ca ngày": "D", "tăng ca đêm": "OD",
                "chủ nhật": "OCN", "ca đêm chủ nhật": "OCND"}
    for m, r in nap_kieu_khoi(da, rows[hang_ngay + 2:], "C", "H", cot_ngay, map_nhan, "Bảng công SK").items():
        them_nv(da, m, ten=ten_sach(o(r, "D")), ma_the=ma(o(r, "B")), ngay_vao=la_ngay(o(r, "E")),
                ngay_nghi=la_ngay(o(r, "F")), bo_phan=ten_sach(o(r, "G")))
    da.tu_ngay, da.den_ngay = min(cot_ngay.values()), max(cot_ngay.values())

    # bảng công nhà máy: cùng bố cục với bảng SK, chỉ khác tên loại giờ
    nm = next((s for s in wb.sheetnames if chuan(s) == "bảng công nhà máy"), None)
    if nm:
        rows = doc_sheet(wb, nm)
        hang_ngay, cot_ngay = hang_co_nhieu_ngay(rows, 10, ci("I"))
        map_nm = {"giờ hành chính": "X", "giờ ban đêm": "OT", "ngoài giờ hành chính": "D",
                  "ngoài giờ hành chính ban đêm": "OD", "ngày nghỉ": "OCN", "ban đêm trong ngày nghỉ": "OCND"}
        for m, r in nap_kieu_khoi(da, rows[hang_ngay + 2:], "C", "H", cot_ngay, map_nm, nm.strip(),
                                  dich=da.moi_bang_nm(nm.strip())).items():
            them_nv(da, m, ten=ten_sach(o(r, "D")), ma_the=ma(o(r, "B")), ngay_vao=la_ngay(o(r, "E")),
                    ngay_nghi=la_ngay(o(r, "F")), bo_phan=ten_sach(o(r, "G")))
        da.mau = "Mẫu B (nhà máy: Bảng công nhà máy; SK: Bảng công SK; đối chiếu Tính phí)"

    if "Tính phí" in wb.sheetnames:
        loai = ["X", "OT", "D", "OD", "OCN", "OCND"]
        da.bang_ke, dong, _ = doc_bang_ke_theo_nv(
            wb, "Tính phí", 8, "C", dict(zip(loai, "HJLNPR")),
            dung_khi=lambda r: chuan(o(r, "A")) == "tổng")
        r6 = doc_sheet(wb, "Tính phí", max_row=6)[5]
        da.don_gia_dv = {c: so(o(r6, col)) for c, col in zip(loai, "IKMOQS") if so(o(r6, col))}
        for m, r in dong.items():
            if m in da.nv:
                them_nv(da, m, ten=ten_sach(o(r, "D")), ngay_vao=la_ngay(o(r, "E")), ngay_nghi=la_ngay(o(r, "F")))
        da.nguon_doi_chieu = da.nguon_don_gia = "Tính phí"
    return da


# --------------------------------------------------------------------------------------
# MẪU C: chấm công 'Tuần' (1 dòng/NV/ngày, giờ đã quy đổi theo 'DM ca'); đối chiếu 'Bảng kê Txx'
# --------------------------------------------------------------------------------------
def nhan_dien_C(sheets):
    return "Tuần" in sheets and "DM ca" in sheets


# cột giờ của mẫu C theo thứ tự HC, CĐ, OT150, OT200 đêm, OT200 CN, OT270, OT300, OT390
LOAI_LAM_C = ["X", "OT", "D", "OD", "OCN", "OCND", "OL", "OLD"]


def doc_mau_C(path, wb) -> DuAn:
    da = DuAn(ma=ma_du_an_tu_file(path, "Dự án C"), file=path, mau="Mẫu C (1 bảng công: Tuần; đối chiếu Bảng kê)", nguon_cc="Tuần")
    rows = doc_sheet(wb, "Tuần")
    hang_td = next(i for i, r in enumerate(rows[:10]) if chuan(o(r, "D")).startswith("mã nhân viên"))
    cot_gio = list(zip("PQRSTUVW", LOAI_LAM_C))
    gap = Counter()
    for r in rows[hang_td + 1:]:
        m, ngay = ma(o(r, "D")), la_ngay(o(r, "A"))
        if not m or not ngay:
            continue
        gap[(m, ngay)] += 1
        # tên bộ phận, nếu trống thì lấy mã bộ phận
        them_nv(da, m, ten=ten_sach(o(r, "E")), bo_phan=ten_sach(o(r, "AK")) or ten_sach(o(r, "AJ")))
        for col, code in cot_gio:
            v = o(r, col)
            if isinstance(v, str) and v.strip().upper() in LOI_CONG_THUC:
                da.o_loi.append((m, "Tuần", f"{col} ({ngay:%d/%m})", v.strip()))
                continue
            h = so(v)
            if h:
                da.cham_cong[(m, ngay)][code] += h
        kh = ten_sach(o(r, "AG"))
        if kh:
            da.ky_hieu_ngay[(m, ngay)] = kh
        da.tu_ngay = min(da.tu_ngay or ngay, ngay)
        da.den_ngay = max(da.den_ngay or ngay, ngay)
    for (m, ngay), n in gap.items():
        if n > 1:
            da.bao(NGHI_NGO, m, "Trùng dòng chấm công", f"Ngày {ngay:%d/%m/%Y} có {n} dòng trong sheet 'Tuần' – giờ đã cộng dồn")

    # 'Tuần' = dữ liệu xuất từ hệ thống nhà máy, giờ quy đổi chuẩn theo mã ca ('DM ca') -> công nhà máy.
    # 'DL công (tính lương)' = cùng dữ liệu đó nhưng SK tự nhập giờ thực tế (W:AD, AI = tổng, AJ = AI - giờ theo
    # mã ca) để tính lương -> công SK. Chỉ dùng khi cùng kỳ với 'Tuần'; không thì dự án coi như 1 bảng công.
    dl = next((s for s in wb.sheetnames if chuan(s).startswith("dl công")), None)
    if dl:
        rows = doc_sheet(wb, dl)
        hang_dl = next((i for i, r in enumerate(rows[:10]) if chuan(o(r, "D")).startswith("mã nhân viên")), None)
        rows = rows[hang_dl + 1:] if hang_dl is not None else []
        if rows and sheet_cung_ky(da, dl, [la_ngay(o(r, "A")) for r in rows], [ma(o(r, "D")) for r in rows if o(r, "D")],
                                  "công tính lương SK"):
            da.cham_cong_nm, da.nguon_nm = da.cham_cong, "Tuần"
            da.cham_cong, da.nguon_cc = defaultdict(lambda: defaultdict(float)), dl
            for r in rows:
                m, ngay = ma(o(r, "D")), la_ngay(o(r, "A"))
                if not m or not ngay:
                    continue
                them_nv(da, m, ten=ten_sach(o(r, "E")))
                for col, code in zip(["W", "X", "Y", "Z", "AA", "AB", "AC", "AD"], LOAI_LAM_C):
                    h = so(o(r, col))
                    if h:
                        da.cham_cong[(m, ngay)][code] += h
            da.mau = "Mẫu C (nhà máy: Tuần; SK: DL công (tính lương); đối chiếu Bảng kê)"

    bk = next((s for s in wb.sheetnames if s.lower().startswith("bảng kê")), None)
    if bk:
        da.bang_ke, dong, da.don_gia_dv = doc_bang_ke_theo_nv(
            wb, bk, 2, "B", dict(zip(["X", "OT", "D", "OD", "OCN", "OCND", "OL", "OLD"], "FHJLNPRT")),
            dict(zip(["X", "OT", "D", "OD", "OCN", "OCND", "OL", "OLD"], "GIKMOQSU")))
        for m, r in dong.items():
            if m in da.nv:
                them_nv(da, m, ngay_vao=la_ngay(o(r, "D")), ngay_nghi=la_ngay(o(r, "E")))
        da.nguon_doi_chieu = da.nguon_don_gia = bk
    return da


# --------------------------------------------------------------------------------------
# MẪU D: danh sách NV + giờ theo NV ở sheet '明细' (bảng kê); chấm công ngày chỉ còn trong cache
# liên kết ngoài 'DS TỔNG' lưu trong file
# --------------------------------------------------------------------------------------
def nhan_dien_D(sheets):
    return any("明细" in s for s in sheets)


def _doc_cache_lien_ket(path, ten_sheet="DS TỔNG"):
    """Trả về list[(tên part, {row: {col: value}})] cho các external link có sheet ten_sheet."""
    kq = []
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if not re.fullmatch(r"xl/externalLinks/externalLink\d+\.xml", n):
                continue
            x = z.read(n).decode("utf8")
            names = re.findall(r'<sheetName val="([^"]*)"', x)
            if ten_sheet not in names:
                continue
            sid = names.index(ten_sheet)
            mblock = re.search(rf'<sheetData sheetId="{sid}"[^>]*>(.*?)</sheetData>', x, re.S)
            if not mblock:
                continue
            data = defaultdict(dict)
            for cm in re.finditer(r'<cell r="([A-Z]+)(\d+)"[^>]*>(?:<v>(.*?)</v>)?</cell>', mblock.group(1), re.S):
                col, row, v = cm.group(1), int(cm.group(2)), cm.group(3)
                if v is not None:
                    v = (v.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
                         .replace("&quot;", '"').replace("&apos;", "'"))
                    data[row][col] = v
            kq.append((n, data))
    return kq


def doc_mau_D(path, wb) -> DuAn:
    da = DuAn(ma=ma_du_an_tu_file(path, "Dự án D"), file=path, mau="Mẫu D (1 bảng công: cache 'DS TỔNG'; đối chiếu 明细)",
              nguon_cc="cache liên kết 'DS TỔNG'")
    # ---- 明细: danh sách NV của hợp đồng + giờ theo từng NV (dùng để đối chiếu)
    mx = next((s for s in wb.sheetnames if "明细" in s and "(" not in s), None) or \
        next(s for s in wb.sheetnames if "明细" in s)
    loai = ["X", "OT", "D", "OD", "OCN", "OCND", "OL", "OLD"]
    # TC ngày / TC đêm có cột thử việc (V 'TGTV 1.5', X 'TV 215%') và chính thức (W, Y) -> cộng cả hai
    da.bang_ke, dong, _ = doc_bang_ke_theo_nv(wb, mx, 5, "B", dict(zip(loai, ["T", "U", ("V", "W"), ("X", "Y"), "Z", "AB", "AD", "AE"])))
    for m, r in dong.items():
        them_nv(da, m, ten=ten_sach(o(r, "C")), bo_phan=" / ".join(x for x in (ten_sach(o(r, "E")), ten_sach(o(r, "G"))) if x),
                chuc_danh=ten_sach(o(r, "I")), ncc=ten_sach(o(r, "N")), ngay_vao=la_ngay(o(r, "J")))
    if dong:
        r5 = next(iter(dong.values()))
        da.don_gia_dv = {c: so(o(r5, col)) for c, col in zip(loai, ["AH", "AJ", "AL", "AN", "AP", "AR", "AT", "AV"])
                         if so(o(r5, col))}
    da.nguon_doi_chieu = da.nguon_don_gia = mx

    th = ky_tu_ten_file(path)
    if th:
        nam, thang = th[1], th[0]
    else:
        nam, thang = dt.date.today().year, dt.date.today().month
        da.bao(LOI, "", "Không xác định được kỳ", "Tên file không có kỳ dạng 'Txx.yyyy' (vd T08.2026); mẫu D cần kỳ "
               f"trong tên file để gán ngày cho chấm công – đang tạm dùng tháng hiện tại {thang:02d}/{nam}. Đổi tên file rồi chạy lại.")
    da.tu_ngay = dt.date(nam, thang, 1)
    da.den_ngay = dt.date(nam + thang // 12, thang % 12 + 1, 1) - dt.timedelta(days=1)

    # ---- chấm công ngày từ cache liên kết ngoài: chọn bản cache khớp giờ '明细' nhiều NV nhất
    def tong_khoi(data, r0):
        return {code: sum(so(data.get(r0 + off, {}).get(get_column_letter(ci("N") + 1 + i))) for i in range(31))
                for off, code in enumerate(["X", "OT", "D", "OD"])}
    tot, diem_tot = None, -1
    for ten_part, data in _doc_cache_lien_ket(path):
        khoi = {}
        for row in sorted(data):
            mm = data[row].get("D", "").strip()
            if mm and re.fullmatch(r"[A-Z]{1,4}\d{3,}", mm):
                khoi.setdefault(mm, row)
        diem = sum(1 for k, r0 in khoi.items() if k in da.bang_ke and
                   all(abs(h - da.bang_ke[k].get(c, 0)) < NGUONG["sai_so_gio"] for c, h in tong_khoi(data, r0).items()))
        if diem > diem_tot:
            tot, diem_tot = (data, khoi), diem
    if not tot:
        da.bao(LOI, "", "Thiếu chấm công ngày", "Không tìm thấy dữ liệu 'DS TỔNG' trong file")
        return da
    data, khoi = tot
    for mm, r0 in khoi.items():
        if mm not in da.nv:
            continue
        for i in range(31):
            try:
                ngay = dt.date(nam, thang, i + 1)
            except ValueError:
                continue
            col = get_column_letter(ci("N") + 1 + i)
            for off, code in enumerate(["X", "OT", "D", "OD"]):
                v = data.get(r0 + off, {}).get(col)
                if v is None:
                    continue
                if re.fullmatch(r"-?[\d.]+", v.strip()):
                    if float(v):
                        da.cham_cong[(mm, ngay)][code] += float(v)
                elif v.strip():
                    da.ky_hieu_ngay[(mm, ngay)] = v.strip()
    khac = len(set(khoi) - set(da.nv))
    da.bao(LUU_Y, "", "Nguồn chấm công", f"Chấm công ngày lấy từ bộ nhớ đệm liên kết ngoài 'DS TỔNG' lưu trong file "
           f"(N = nghỉ, TD = nghỉ tự do); chỉ lấy {len(da.nv)} NV có trong '{mx}'"
           + (f", bỏ qua {khac} mã của hợp đồng khác trong cache." if khac else "."))
    return da


# --------------------------------------------------------------------------------------
# MẪU E: chấm công 'Công (lương)' (7 dòng/NV theo hệ số); đối chiếu 'Công' (bảng công phía khách hàng)
# --------------------------------------------------------------------------------------
def nhan_dien_E(sheets):
    return "Công (lương)" in sheets


HE_SO_E = {1.0: "X", 1.3: "OT", 1.5: "D", 2.0: "OCN", 2.3: "OCND", 3.0: "OD", 3.3: "OLD"}


def _doc_cong_he_so(da, rows, sheet, cham_cong):
    """Sheet công dạng 'hệ số' (mẫu E): dòng 2 là ngày, mỗi NV nhiều dòng, cột F = hệ số. Trả về {mã: dòng đầu}."""
    cot_ngay = {get_column_letter(j + 1): v.date() for j, v in enumerate(rows[1]) if isinstance(v, dt.datetime)}
    hien_tai, dau = None, {}
    for r in rows[4:]:
        if chuan(o(r, "A")) == "total":
            break
        m = ma(o(r, "B"))
        if m:
            if m in dau and m != hien_tai and cham_cong is da.cham_cong:
                da.trung_ma[m] = da.trung_ma.get(m, 1) + 1
            dau.setdefault(m, r)
            hien_tai = m
        if hien_tai is None or o(r, "F") in (None, ""):
            continue
        code = HE_SO_E.get(round(so(o(r, "F")), 2))
        if not code:
            if so(o(r, "F")) and any(so(o(r, c)) for c in cot_ngay):
                da.bao(NGHI_NGO, hien_tai, "Nhãn giờ lạ", f"'{sheet}': hệ số '{o(r, 'F')}' không thuộc bảng hệ số đã biết – bỏ qua dòng này")
            continue
        for col, ngay in cot_ngay.items():
            v = o(r, col)
            if isinstance(v, str) and v.strip().upper() in LOI_CONG_THUC:
                if cham_cong is da.cham_cong:
                    da.o_loi.append((hien_tai, sheet, f"{col} ({ngay:%d/%m})", v.strip()))
                continue
            h = so(v)
            if h:
                cham_cong[(hien_tai, ngay)][code] += h
    return cot_ngay, dau


def doc_mau_E(path, wb) -> DuAn:
    # 'Công (lương)' = công nhà máy: số liệu chết trích từ phần mềm chấm công nhà máy.
    # 'Công' = công SK: SK dùng công thức SUMIFS trên dữ liệu gốc 'DL công' để kiểm tra lại công nhà máy.
    da = DuAn(ma=ma_du_an_tu_file(path, "Dự án E"), file=path,
              mau="Mẫu E (nhà máy: Công (lương); SK: Công)", nguon_cc="Công",
              loai_mau=list(HE_SO_E.values()))     # mỗi NV 7 dòng hệ số như file gốc
    rows = doc_sheet(wb, "Công (lương)")
    cot_ngay, dau = _doc_cong_he_so(da, rows, "Công (lương)", da.moi_bang_nm("Công (lương)"))
    da.tu_ngay, da.den_ngay = min(cot_ngay.values()), max(cot_ngay.values())
    for m, r in dau.items():
        them_nv(da, m, ten=ten_sach(o(r, "C")), bo_phan=ten_sach(o(r, "D")), ngay_vao=la_ngay(o(r, "E")))
    if "Công" in wb.sheetnames:
        _, dau = _doc_cong_he_so(da, doc_sheet(wb, "Công"), "Công", da.cham_cong)
        for m, r in dau.items():
            them_nv(da, m, ten=ten_sach(o(r, "C")), bo_phan=ten_sach(o(r, "D")), ngay_vao=la_ngay(o(r, "E")))
    else:                                   # chỉ có 1 bảng công: dùng chung cho nhà máy và SK
        da.cham_cong, da.cham_cong_nm, da.nguon_cc, da.nguon_nm = da.cham_cong_nm, None, "Công (lương)", ""
        da.mau = "Mẫu E (1 bảng công: Công (lương))"
    # Đơn giá ở dòng 3 sheet 'Công (lương)' (cao hơn đơn giá lương -> nhiều khả năng là đơn giá dịch vụ)
    r3 = rows[2] if len(rows) > 2 else ()
    da.don_gia_dv = {c: so(o(r3, col)) for c, col in zip(["X", "OT", "D", "OCN", "OCND", "OD", "OLD"],
                                                        ["AL", "AM", "AN", "AO", "AP", "AQ", "AR"]) if so(o(r3, col))}
    da.nguon_don_gia = "Công (lương) dòng 3 (chưa xác nhận là đơn giá dịch vụ)"
    # danh sách NV: tên chuẩn (viết hoa đầu chữ), ngày vào, ngày nghỉ
    for sh, cot_nghi in (("DS CN", "F"), ("DS CN (nghỉ)", "G")):
        if sh not in wb.sheetnames:
            continue
        for r in doc_sheet(wb, sh)[1:]:
            m = ma(o(r, "B"))
            if m not in da.nv:
                continue
            nv = da.nv[m]
            if ten_sach(o(r, "D")) and chuan(o(r, "D")) == chuan(nv.ten):
                nv.ten = ten_sach(o(r, "D"))
            them_nv(da, m, ngay_vao=la_ngay(o(r, "E")), ngay_nghi=la_ngay(o(r, cot_nghi)))
    return da


ADAPTERS = [
    ("A", nhan_dien_A, doc_mau_A),
    ("B", nhan_dien_B, doc_mau_B),
    ("C", nhan_dien_C, doc_mau_C),
    ("D", nhan_dien_D, doc_mau_D),
    ("E", nhan_dien_E, doc_mau_E),
]


def doc_file(path, log=print) -> DuAn:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for ten, nd, doc in ADAPTERS:
            if nd(wb.sheetnames):
                log(f"  -> nhận diện biểu mẫu {ten}")
                da = doc(path, wb)
                break
        else:
            raise ValueError(f"Không nhận diện được biểu mẫu của file {os.path.basename(path)} "
                             f"(sheets: {wb.sheetnames})")
    except ValueError:
        raise
    except Exception as e:      # StopIteration (không thấy dòng tiêu đề), sheet rỗng, ...
        raise ValueError(f"File {os.path.basename(path)} nhận diện là mẫu {ten} nhưng đọc không được "
                         f"({type(e).__name__}: {e or 'không tìm thấy dòng tiêu đề/dữ liệu'}). "
                         "Có thể biểu mẫu đã bị đổi cấu trúc.") from e
    finally:
        wb.close()
    hoan_thien(da)
    kiem_tra_du_an(da)
    return da


def hoan_thien(da: DuAn):
    """Xác định kỳ lương từ tên file / kỳ công."""
    th = ky_tu_ten_file(da.file)
    if not da.tu_ngay:
        ngay = sorted({k[1] for k in da.cham_cong} | {k[1] for k in da.ky_hieu_ngay})
        if ngay:
            da.tu_ngay, da.den_ngay = ngay[0], ngay[-1]
        elif th:
            da.tu_ngay = dt.date(th[1], th[0], 1)
            da.den_ngay = dt.date(th[1] + th[0] // 12, th[0] % 12 + 1, 1) - dt.timedelta(days=1)
    if not th and da.den_ngay:
        th = (da.den_ngay.month, da.den_ngay.year)     # kỳ 26/07 - 25/08 => T08
    da.ky = f"T{th[0]:02d}/{th[1]}" if th else ""


def doi_chieu_nm_sk(da: DuAn) -> dict | None:
    """So công nhà máy với công SK theo từng NV, từng loại giờ. None nếu dự án chỉ có 1 bảng công.
    Trả về {gio_nm, gio_sk, nv_lech: {mã: {code: (nhà máy, SK)}}, chi_nm, chi_sk, khop}."""
    if not da.hai_bang:
        return None
    sai = NGUONG["sai_so_gio"]
    nm, sk = da.tong_gio_nv(da.cham_cong_nm), da.tong_gio_nv()
    kq = {"gio_nm": sum(sum(g.get(c, 0) for c in LOAI_LAM) for g in nm.values()),
          "gio_sk": sum(sum(g.get(c, 0) for c in LOAI_LAM) for g in sk.values()),
          "nv_lech": {}, "chi_nm": [], "chi_sk": []}
    for m in da.nv:
        a, b = nm.get(m, {}), sk.get(m, {})
        ta, tb = sum(a.values()), sum(b.values())
        if ta > sai and tb <= sai:
            kq["chi_nm"].append(m)
        elif tb > sai and ta <= sai:
            kq["chi_sk"].append(m)
        else:
            lech = {c: (a.get(c, 0), b.get(c, 0)) for c in LOAI_LAM if abs(a.get(c, 0) - b.get(c, 0)) > sai}
            if lech:
                kq["nv_lech"][m] = lech
    kq["khop"] = not (kq["nv_lech"] or kq["chi_nm"] or kq["chi_sk"])
    return kq


def doi_chieu(da: DuAn) -> dict | None:
    """So công NHÀ MÁY với bảng kê (bảng tính tiền với nhà máy). None nếu file không có bảng kê.
    Trả về {nguon, gio_cc, gio_bk, nv_lech: {mã: {code: (công, bảng kê)}}, chi_cong, chi_bk, loai_lech}."""
    if da.bang_ke is None and da.bang_ke_tong is None:
        return None
    sai = NGUONG["sai_so_gio"]
    tong = da.tong_gio_nv(da.cc_nm)
    gio_cc = sum(sum(g.get(c, 0) for c in LOAI_LAM) for g in tong.values())
    kq = {"nguon": da.nguon_doi_chieu, "gio_cc": gio_cc, "nv_lech": {}, "chi_cong": [], "chi_bk": [], "loai_lech": {}}
    if da.bang_ke is not None:
        bk = da.bang_ke
        kq["gio_bk"] = sum(sum(g.get(c, 0) for c in LOAI_LAM) for g in bk.values())
        for m in da.nv:
            g = tong.get(m, {})
            if m not in bk:
                if sum(g.values()) > sai:
                    kq["chi_cong"].append(m)
                continue
            lech = {c: (g.get(c, 0), bk[m].get(c, 0)) for c in LOAI_LAM if abs(g.get(c, 0) - bk[m].get(c, 0)) > sai}
            if lech:
                kq["nv_lech"][m] = lech
        kq["chi_bk"] = [m for m, g in bk.items() if m not in da.nv and sum(g.values()) > sai]
    else:
        tong_loai = defaultdict(float)
        for g in tong.values():
            for c, h in g.items():
                tong_loai[c] += h
        kq["gio_bk"] = sum(da.bang_ke_tong.get(c, 0) for c in LOAI_LAM)
        kq["loai_lech"] = {c: (tong_loai.get(c, 0), da.bang_ke_tong.get(c, 0)) for c in LOAI_LAM
                           if abs(tong_loai.get(c, 0) - da.bang_ke_tong.get(c, 0)) > sai}
    kq["khop"] = not (kq["nv_lech"] or kq["chi_cong"] or kq["chi_bk"] or kq["loai_lech"])
    return kq


# --------------------------------------------------------------------------------------
# KIỂM TRA LỖI / NGHI NGỜ
# --------------------------------------------------------------------------------------
def _rut_gon(ds, n=15):
    ds = list(ds)
    return ", ".join(ds[:n]) + (f" ... (+{len(ds) - n})" if len(ds) > n else "")


def _mau_ma(m):
    # chỉ so cấu trúc chữ/số (HMT02 và HMT474 cùng dạng), không so độ dài dãy số
    return re.sub(r"9+", "9", re.sub(r"[A-Za-z]+", "A", re.sub(r"\d", "9", m)))


def _ten_loai(c):
    return TEN_LOAI_TUAN.get(c, c)


def kiem_tra_du_an(da: DuAn):
    """Các quy tắc phát hiện sai sót / bất thường trong 1 hợp đồng (chỉ dựa trên chấm công và bảng kê)."""
    # 1. File / cấu trúc
    if not da.nv:
        da.bao(LOI, "", "Bảng chấm công trống", f"Không đọc được nhân viên nào trong '{da.nguon_cc}' - kiểm tra lại biểu mẫu")
    th = ky_tu_ten_file(da.file)
    if th and da.tu_ngay and da.den_ngay and th[0] not in (da.tu_ngay.month, da.den_ngay.month):
        da.bao(NGHI_NGO, "", "Kỳ lương không khớp", f"Tên file ghi T{th[0]:02d}/{th[1]} nhưng bảng công là "
               f"{da.tu_ngay:%d/%m/%Y} – {da.den_ngay:%d/%m/%Y}")
    if da.tu_ngay and da.den_ngay and (da.den_ngay - da.tu_ngay).days + 1 > 31:
        da.bao(LOI, "", "Kỳ công dài hơn 31 ngày",
               f"Kỳ {da.tu_ngay:%d/%m/%Y} – {da.den_ngay:%d/%m/%Y} dài {(da.den_ngay - da.tu_ngay).days + 1} ngày; tab "
               "bang_cham_cong_thang chỉ có 31 cột ngày nên công sau ngày thứ 31 không hiện trong lưới. Kiểm tra ngày gõ nhầm trong bảng công.")
    if da.nv and not da.cham_cong:
        da.bao(LOI, "", "Không có giờ công", f"Có {len(da.nv)} NV nhưng mọi ô giờ đều trống. Nếu ô giờ là công thức, có thể file "
               "được lưu bằng công cụ không tính công thức (giá trị chưa được lưu): mở file bằng Excel, bấm Lưu, rồi chạy lại.")
    theo_nv = defaultdict(list)
    for m, sh, cell, v in da.o_loi:
        theo_nv[m].append(f"{sh}!{cell}={v}")
    for m, ds in theo_nv.items():
        da.bao(LOI, m, "Ô lỗi công thức", "Ô bị lỗi, giá trị được tính là 0: " + "; ".join(ds[:8]) + (" ..." if len(ds) > 8 else ""))
    for m, n in da.trung_ma.items():
        da.bao(LOI, m, "Trùng mã NV", f"Mã xuất hiện ở {n} khối NV khác nhau trong '{da.nguon_cc}' – giờ đã bị cộng dồn vào 1 NV")

    # 2. Từng nhân viên
    tong = da.tong_gio_nv()
    mau_pho_bien = Counter(_mau_ma(m) for m in da.nv)
    mau_chinh = mau_pho_bien.most_common(1)[0][0] if mau_pho_bien else None
    vd_ma = next((m for m in da.nv if _mau_ma(m) == mau_chinh), "")
    vuot_tc, khong_gio, chu_han = [], [], []
    trung_ten = defaultdict(list)
    for m, nv in da.nv.items():
        g = tong.get(m, {})
        if not nv.ten:
            da.bao(LOI, m, "Thiếu tên NV", f"Không có họ tên trong '{da.nguon_cc}'")
        else:
            trung_ten[(chuan(nv.ten), nv.ngay_vao)].append(m)
            if co_chu_han(nv.ten):
                chu_han.append(f"{m} {nv.ten}")
        if mau_chinh and mau_pho_bien[mau_chinh] >= 5 and _mau_ma(m) != mau_chinh:
            da.bao(NGHI_NGO, m, "Mã NV khác định dạng", f"Mã '{m}' khác dạng của các mã khác trong hợp đồng (vd '{vd_ma}') – có thể gõ thừa khoảng trắng/ký tự")
        if sum(g.values()) <= 0:
            khong_gio.append(m)
        for c, h in g.items():
            if h < 0:
                da.bao(LOI, m, "Giờ âm", f"{_ten_loai(c)}: tổng {h:g} giờ")
        tc = sum(g.get(c, 0) for c in LOAI_TANG_CA)
        if tc > NGUONG["tang_ca_thang"]:
            vuot_tc.append((tc, m))
        if nv.ngay_vao and nv.ngay_nghi and nv.ngay_nghi < nv.ngay_vao:
            da.bao(LOI, m, "Ngày nghỉ trước ngày vào", f"Ngày vào {nv.ngay_vao:%d/%m/%Y}, ngày nghỉ {nv.ngay_nghi:%d/%m/%Y}")
        if nv.ngay_vao and da.den_ngay and nv.ngay_vao > da.den_ngay and sum(g.values()):
            da.bao(NGHI_NGO, m, "Ngày vào sau kỳ công", f"Ngày vào {nv.ngay_vao:%d/%m/%Y} nhưng có giờ trong kỳ đến {da.den_ngay:%d/%m/%Y}")
    if khong_gio and da.cham_cong:
        da.bao(LUU_Y, "", "NV không có giờ công", f"{len(khong_gio)} NV có tên trong bảng công nhưng không có giờ nào trong kỳ "
               "(vẫn giữ trong file chuẩn): " + _rut_gon(khong_gio, 20))
    if chu_han:
        da.bao(LUU_Y, "", "Tên chữ Hán", f"{len(chu_han)} NV có tên chữ Hán trong nguồn, cần sửa tên tiếng Việt: " + _rut_gon(chu_han, 10))
    if vuot_tc:
        vuot_tc.sort(reverse=True)
        da.bao(LUU_Y, "", "Làm thêm vượt 40 giờ/tháng",
               f"{len(vuot_tc)} NV có tổng giờ tăng ca > {NGUONG['tang_ca_thang']} giờ/kỳ (BLLĐ 2019 Điều 107). "
               f"Cao nhất: " + _rut_gon(f"{m} ({h:g}h)" for h, m in vuot_tc[:10]))
    for (ten, nv_vao), ds in trung_ten.items():
        if len(ds) > 1:
            da.bao(NGHI_NGO, ds[0], "Có thể trùng người", f"{len(ds)} mã cùng họ tên '{da.nv[ds[0]].ten}'"
                   + (f" và cùng ngày vào {nv_vao:%d/%m/%Y}" if nv_vao else "") + ": " + ", ".join(ds))

    # 3. Chấm công theo ngày
    co_tang_ca_cn = any(g.get("OCN") or g.get("OCND") for g in tong.values())
    ngoai_ky, qua_gio, sau_nghi, truoc_vao, hc_cn = (defaultdict(list) for _ in range(5))
    for (m, ngay), g in da.cham_cong.items():
        if m not in da.nv:
            continue
        nv, tg = da.nv[m], sum(g.values())
        if da.tu_ngay and tg and (not (da.tu_ngay <= ngay <= da.den_ngay) or (ngay - da.tu_ngay).days >= 31):
            ngoai_ky[m].append(ngay)
        if tg > NGUONG["gio_ngay_nghi_ngo"]:
            qua_gio[m].append((ngay, tg))
        if nv.ngay_nghi and ngay > nv.ngay_nghi and tg:
            sau_nghi[m].append(ngay)
        if nv.ngay_vao and ngay < nv.ngay_vao and tg:
            truoc_vao[m].append(ngay)
        if co_tang_ca_cn and ngay.weekday() == 6 and any(g.get(c) for c in LOAI_CHINH):
            hc_cn[m].append(ngay)
    for m, ds in ngoai_ky.items():
        da.bao(LOI, m, "Công ngoài kỳ", f"{len(ds)} ngày có công nằm ngoài kỳ / ngoài 31 cột ngày: "
               + _rut_gon(f"{n:%d/%m/%Y}" for n in sorted(ds)))
    for m, ds in qua_gio.items():
        nang = [x for x in ds if x[1] > NGUONG["gio_ngay_loi"]]
        da.bao(LOI if nang else NGHI_NGO, m, "Giờ làm trong ngày quá cao",
               f"{len(ds)} ngày > {NGUONG['gio_ngay_nghi_ngo']} giờ: " + _rut_gon(f"{n:%d/%m} ({h:g}h)" for n, h in sorted(ds)))
    for m, ds in sau_nghi.items():
        da.bao(NGHI_NGO, m, "Có công sau ngày nghỉ việc", f"Nghỉ việc {da.nv[m].ngay_nghi:%d/%m/%Y} nhưng còn công {len(ds)} ngày: "
               + _rut_gon(f"{n:%d/%m}" for n in sorted(ds)))
    for m, ds in truoc_vao.items():
        da.bao(NGHI_NGO, m, "Có công trước ngày vào", f"Vào làm {da.nv[m].ngay_vao:%d/%m/%Y} nhưng có công {len(ds)} ngày trước đó: "
               + _rut_gon(f"{n:%d/%m}" for n in sorted(ds)))
    for m, ds in hc_cn.items():
        da.bao(NGHI_NGO, m, "Giờ thường ghi vào Chủ nhật", f"{len(ds)} ngày Chủ nhật có giờ HC/tăng ca ngày thường "
               "(hợp đồng có tăng ca CN riêng): " + _rut_gon(f"{n:%d/%m}" for n in sorted(ds)))

    # 4. Đối chiếu công nhà máy với công SK
    ns = doi_chieu_nm_sk(da)
    if ns:
        for m, lech in ns["nv_lech"].items():
            da.bao(NGHI_NGO, m, "Lệch công nhà máy / SK", f"'{da.nguon_nm}' / '{da.nguon_cc}': " + "; ".join(
                f"{_ten_loai(c)}: nhà máy {a:g} / SK {b:g}" for c, (a, b) in lech.items()))
        if ns["chi_nm"]:
            da.bao(NGHI_NGO, "", "Có công nhà máy nhưng không có công SK",
                   f"{len(ns['chi_nm'])} NV có giờ trong '{da.nguon_nm}' nhưng không có giờ trong '{da.nguon_cc}': "
                   + _rut_gon(ns["chi_nm"], 20))
        if ns["chi_sk"]:
            da.bao(NGHI_NGO, "", "Có công SK nhưng không có công nhà máy",
                   f"{len(ns['chi_sk'])} NV có giờ trong '{da.nguon_cc}' nhưng không có giờ trong '{da.nguon_nm}': "
                   + _rut_gon(ns["chi_sk"], 20))

    # 5. Đối chiếu công nhà máy với bảng kê
    dc = doi_chieu(da)
    if dc is None:
        da.bao(LUU_Y, "", "Không có bảng kê", "File không có bảng kê để đối chiếu giờ chấm công")
        return
    for m, lech in dc["nv_lech"].items():
        da.bao(NGHI_NGO, m, "Lệch giờ chấm công / bảng kê", f"Công nhà máy so với '{dc['nguon']}': " + "; ".join(
            f"{_ten_loai(c)}: công {a:g} / bảng kê {b:g}" for c, (a, b) in lech.items()))
    if dc["chi_cong"]:
        da.bao(NGHI_NGO, "", "Có công nhưng không có trong bảng kê",
               f"{len(dc['chi_cong'])} NV có giờ trong công nhà máy nhưng không có trong '{dc['nguon']}': " + _rut_gon(dc["chi_cong"], 20))
    if dc["chi_bk"]:
        da.bao(NGHI_NGO, "", "Có trong bảng kê nhưng không có chấm công",
               f"{len(dc['chi_bk'])} mã có giờ trong '{dc['nguon']}' nhưng không có trong bảng công: " + _rut_gon(dc["chi_bk"], 20))
    if dc["loai_lech"]:
        da.bao(NGHI_NGO, "", "Lệch tổng giờ với bảng kê", f"So với '{dc['nguon']}': " + "; ".join(
            f"{_ten_loai(c)}: công {a:g} / bảng kê {b:g}" for c, (a, b) in dc["loai_lech"].items()))


def kiem_tra_cheo(ds: list[DuAn]):
    """Kiểm tra chéo giữa các hợp đồng (chạy sau khi đã đọc hết file)."""
    theo_ma = defaultdict(list)
    for da in ds:
        for m in da.nv:
            theo_ma[m].append(da)
    for m, das in theo_ma.items():
        if len({x.ma for x in das}) > 1:
            for da in das:
                da.bao(NGHI_NGO, m, "Trùng mã giữa các hợp đồng",
                       f"Mã {m} có trong {', '.join(sorted({x.ma for x in das}))} - kiểm tra có phải cùng một người")
    dem_ky = Counter(da.ky for da in ds if da.ky)
    if len(dem_ky) > 1:
        ky_chinh = dem_ky.most_common(1)[0][0]
        for da in ds:
            if da.ky and da.ky != ky_chinh:
                da.bao(LUU_Y, "", "Khác kỳ lương", f"Hợp đồng này kỳ {da.ky}, các hợp đồng khác phần lớn kỳ {ky_chinh}. "
                       "Kiểm tra có chọn nhầm file tháng khác không.")
    for da in ds:
        da.van_de.sort(key=lambda x: (THU_TU_MUC[x[0]], x[2], x[1]))


def dem_muc(da: DuAn) -> dict:
    dem = {LOI: 0, NGHI_NGO: 0, LUU_Y: 0}
    for x in da.van_de:
        dem[x[0]] += 1
    return dem


# --------------------------------------------------------------------------------------
# GHI FILE CHUẨN
# --------------------------------------------------------------------------------------
FONT = "Arial"
F_TD = Font(name=FONT, bold=True, size=10)
F_TH = Font(name=FONT, size=10)
F_GHI = Font(name=FONT, size=9, italic=True, color="666666")
NEN_TD = PatternFill("solid", fgColor="D9E1F2")
NEN_PHU = PatternFill("solid", fgColor="FFF2CC")
NEN_DO = PatternFill("solid", fgColor="F8CBAD")
NEN_MUC = {"Lỗi": PatternFill("solid", fgColor="F4B6A6"), "Nghi ngờ": PatternFill("solid", fgColor="FCE4B0"),
           "Lưu ý": PatternFill("solid", fgColor="DDEBF7")}
VIEN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
CANH_GIUA = Alignment(horizontal="center", vertical="center", wrap_text=True)
SO = "#,##0;-#,##0;-"
GIO = "#,##0.##;-#,##0.##;-"
THU = {0: "Thứ 2", 1: "Thứ 3", 2: "Thứ 4", 3: "Thứ 5", 4: "Thứ 6", 5: "Thứ 7", 6: "Chủ nhật"}


def _tieu_de(ws, hang, ds, nen=NEN_TD, bat_dau=1):
    for j, t in enumerate(ds, start=bat_dau):
        c = ws.cell(hang, j, t)
        c.font, c.fill, c.alignment, c.border = F_TD, nen, CANH_GIUA, VIEN


def _o(ws, r, c, v, fmt=None, font=F_TH):
    cell = ws.cell(r, c, v)
    # gán font/viền từng ô rất chậm ở openpyxl (băm lại style mỗi lần); chép mảng style của ô đầu tiên
    # cùng kiểu. Mảng style trỏ vào danh sách style của workbook nên bộ đệm lưu theo từng workbook.
    # Chỉ áp cho ô chưa có style, để không xóa nền/căn lề đã gán trước đó.
    dem = ws.parent.__dict__.setdefault("_kieu_o", {})
    mau = dem.get((id(font), fmt))
    if mau is not None and not cell.has_style:
        cell._style = copy(mau)
        return cell
    moi = not cell.has_style
    cell.font, cell.border = font, VIEN
    if fmt:
        cell.number_format = fmt
    if moi:
        dem[(id(font), fmt)] = copy(cell._style)
    return cell


def _khoi_cham_cong(ws, r, stt, da, ngay, ky, tieu_de, cot_phu, ten_phu, gt_phu, gop=True,
                    ten_cot_e="Tổng công ca ngày", nhan=TEN_LOAI, cc=None, khac=None):
    """Ghi 1 khối chấm công theo template (hàng thứ, hàng tiêu đề có ngày, 1 hàng trống, rồi dữ liệu từ hàng 4).
    ngay: các ngày ứng với cột G trở đi. ky=None (bảng tháng): cột F = toàn bộ giờ của NV trong kỳ;
    ky=(từ, đến) (bảng tuần): cột F = giờ trong khối, ngày ngoài kỳ tô nhạt ở hàng tiêu đề.
    gop: gộp ô Mã NV/Tên theo NV (template tháng có gộp, template tuần không).
    ten_cot_e, nhan: tiêu đề cột E và tên loại giờ (template tuần ghi khác template tháng).
    cc: bảng công cần ghi (mặc định công SK). khac: bảng công để so, ô khác giá trị được tô đỏ.
    Trả về (hàng tiếp theo sau khối, stt)."""
    cc_nv = defaultdict(dict)                    # mã NV -> {ngày: giờ theo loại}
    for (m, n), g in (da.cham_cong if cc is None else cc).items():
        cc_nv[m][n] = g
    khac_nv = defaultdict(dict)
    for (m, n), g in (khac or {}).items():
        khac_nv[m][n] = g
    cot_cua_ngay = {n: 7 + i for i, n in enumerate(ngay)}
    if tieu_de:
        ws.cell(r, 1, tieu_de).font = Font(name=FONT, bold=True, size=10, color="C00000")
    for n, col in cot_cua_ngay.items():
        ws.cell(r, col, THU[n.weekday()]).font = F_TH
    r += 1
    _tieu_de(ws, r, ["STT", "Mã theo dõi", "Mã NV", "Tên nhân viên", ten_cot_e, "TỔNG CỘNG"])
    for n, col in cot_cua_ngay.items():
        c = ws.cell(r, col, dt.datetime.combine(n, dt.time()))
        c.number_format = "yyyy-mm-dd"
        ngoai = ky and not (ky[0] <= n <= ky[1])
        c.font, c.fill, c.alignment, c.border = F_TD, NEN_PHU if ngoai else NEN_TD, CANH_GIUA, VIEN
    _tieu_de(ws, r, ten_phu, nen=NEN_PHU, bat_dau=cot_phu)
    r += 2      # bỏ trống 1 hàng như template (dữ liệu bắt đầu ở hàng thứ 4 của khối)
    for nv in da.nv.values():
        ngay_nv = cc_nv.get(nv.ma_nv, {})
        ngay_khac = khac_nv.get(nv.ma_nv, {})
        co = {c for d in (ngay_nv, ngay_khac) for g in d.values() for c, h in g.items() if h}
        loai = LOAI_CHINH + [c for c in THU_TU_GIO_PHU if c in co or c in da.loai_mau]
        stt += 1
        r0 = r
        for code in loai:
            _o(ws, r, 1, stt if r == r0 else None)
            _o(ws, r, 2, nv.ma_nv, "@")
            _o(ws, r, 3, nv.ma_nv if r == r0 else None, "@")
            _o(ws, r, 4, nv.ten if r == r0 else None)
            _o(ws, r, 5, nhan[code])
            tong = 0
            for n, g in (ngay_nv.items() if ky is None else ((n, ngay_nv[n]) for n in ngay if n in ngay_nv)):
                h = g.get(code, 0)
                if h:
                    if n in cot_cua_ngay:
                        _o(ws, r, cot_cua_ngay[n], h, GIO)
                    tong += h
            for col in cot_cua_ngay.values():
                if ws.cell(r, col).value is None:
                    _o(ws, r, col, None)
            _o(ws, r, 6, tong, GIO)
            if khac is not None:
                for n, col in cot_cua_ngay.items():
                    if abs(ngay_nv.get(n, {}).get(code, 0) - ngay_khac.get(n, {}).get(code, 0)) > NGUONG["sai_so_gio"]:
                        ws.cell(r, col).fill = NEN_DO
                # cột F: bảng tháng cộng mọi ngày của NV, bảng tuần chỉ cộng các ngày của khối (như tong ở trên)
                tong_khac = sum(g.get(code, 0) for n, g in ngay_khac.items() if ky is None or n in cot_cua_ngay)
                if abs(tong - tong_khac) > NGUONG["sai_so_gio"]:
                    ws.cell(r, 6).fill = NEN_DO
            for j, v in enumerate(gt_phu, start=cot_phu):
                _o(ws, r, j, v, "dd/mm/yyyy" if isinstance(v, dt.date) else None)
            r += 1
        if gop and r - r0 > 1:
            ws.merge_cells(start_row=r0, end_row=r - 1, start_column=3, end_column=3)
            ws.merge_cells(start_row=r0, end_row=r - 1, start_column=4, end_column=4)
            for c in (3, 4):
                ws.cell(r0, c).alignment = Alignment(vertical="center", wrap_text=True)
    return r + 2, stt      # cách 2 hàng trước khối tiếp theo


TAB_CHAM_CONG = [
    ("bang_cham_cong_thang_nha_may", lambda da: da.cc_nm, "thang"),     # doanh thu SK thu của nhà máy
    ("bang_cham_cong_tuan_nha_may", lambda da: da.cc_nm, "tuan"),
    ("bang_cham_cong_thang_SK", lambda da: da.cham_cong, "thang"),      # chi phí SK trả người lao động
    ("bang_cham_cong_tuan_SK", lambda da: da.cham_cong, "tuan"),
]


def ten_tab_du_an(da: DuAn, da_co):
    """Tên tab riêng của dự án: tên hợp đồng, bỏ ký tự Excel cấm, tối đa 31 ký tự, không trùng tab khác."""
    goc = re.sub(r"[\[\]:*?/\\']", "", da.ma).strip()[:31] or "Du an"
    ten, i = goc, 2
    while ten.lower() in {x.lower() for x in da_co}:
        ten = f"{goc[:27]} ({i})"
        i += 1
    return ten


def _tieu_de_hd(da):
    return (f"HỢP ĐỒNG: {da.ma} | Kỳ {da.ky} | {da.tu_ngay:%d/%m/%Y} – {da.den_ngay:%d/%m/%Y}"
            if da.tu_ngay else f"HỢP ĐỒNG: {da.ma}")


def _ghi_thang(ws, r, stt, da, cc, tieu_de, khac=None):
    """1 khối bảng tháng giống hệt template: hàng thứ + hàng tiêu đề có NGÀY THẬT theo đúng kỳ công của bảng
    chấm công gốc (vd 26/07 -> 25/08), cột G = ngày đầu kỳ, H = ngày thứ 2..."""
    ngay_ky = []
    if da.tu_ngay:
        n = da.tu_ngay
        while n <= da.den_ngay and len(ngay_ky) < 31:
            ngay_ky.append(n)
            n += dt.timedelta(days=1)
    return _khoi_cham_cong(ws, r, stt, da, ngay_ky, None, tieu_de, 39,
                           ["Hợp đồng / Dự án", "Kỳ lương", "Từ ngày", "Đến ngày"],
                           [da.ma, da.ky, da.tu_ngay, da.den_ngay], cc=cc, khac=khac)


def _ghi_tuan(ws, r, da, cc, co_tieu_de, tieu_de=None, khac=None):
    """Mỗi tuần (Thứ 2 -> Chủ nhật) của kỳ công là 1 khối giống template: G = Thứ 2 ... M = Chủ nhật.
    Tuần đầu/cuối kỳ có thể chứa ngày ngoài kỳ: vẫn hiện ngày (tô nhạt) để cột luôn đúng thứ."""
    if not da.tu_ngay:
        return r
    dau = da.tu_ngay - dt.timedelta(days=da.tu_ngay.weekday())
    so_tuan = 0
    while dau <= da.den_ngay:
        so_tuan += 1
        tuan = [dau + dt.timedelta(days=i) for i in range(7)]
        # file 1 hợp đồng để trống A1 như template (tuần đã thể hiện ở hàng ngày và cột O:S)
        td = None
        if co_tieu_de:
            td = (f"{tieu_de} | " if tieu_de else f"HỢP ĐỒNG: {da.ma} | Kỳ {da.ky} | ") +                  f"TUẦN {so_tuan}: {max(dau, da.tu_ngay):%d/%m/%Y} – {min(tuan[-1], da.den_ngay):%d/%m/%Y}"
        r, _ = _khoi_cham_cong(ws, r, 0, da, tuan, (da.tu_ngay, da.den_ngay), td, 15,
                               ["Hợp đồng / Dự án", "Kỳ lương", "Tuần", "Từ ngày", "Đến ngày"],
                               [da.ma, da.ky, so_tuan, tuan[0], tuan[-1]], gop=False,
                               ten_cot_e="Loại công", nhan=TEN_LOAI_TUAN, cc=cc, khac=khac)
        dau += dt.timedelta(days=7)
    return r


def _dinh_dang_thang(ws):
    for col, w in {"A": 5, "B": 12, "C": 12, "D": 24, "E": 20, "F": 10}.items():
        ws.column_dimensions[col].width = w
    for k in range(7, 38):
        ws.column_dimensions[get_column_letter(k)].width = 10.5
    for col, w in {"AM": 9, "AN": 10, "AO": 11, "AP": 11}.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "G1"


def _dinh_dang_tuan(ws):
    for col, w in {"A": 5, "B": 12, "C": 12, "D": 24, "E": 20, "F": 10, "N": 3,
                   "O": 9, "P": 10, "Q": 6, "R": 11, "S": 11}.items():
        ws.column_dimensions[col].width = w
    for k in range(7, 14):
        ws.column_dimensions[get_column_letter(k)].width = 10.5
    ws.freeze_panes = "G1"


def ghi_file_chuan(ds: list[DuAn], out_path, log=print, file_loi=()):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ================= DM =================
    # A–F: danh mục loại giờ (8 dòng đầu đúng template). Đơn giá lương để trống (không đọc bảng lương);
    # đơn giá dịch vụ lấy từ bảng kê khi file chỉ có 1 hợp đồng.
    ws = wb.create_sheet("DM")
    _tieu_de(ws, 1, ["STT", "Ký hiệu", "Nội dung", "Đơn vị tính", "Đơn giá lương", "Đơn giá dịch vụ"])
    for i, (code, ten) in enumerate(LOAI_GIO, start=1):
        _o(ws, i + 1, 1, i); _o(ws, i + 1, 2, code); _o(ws, i + 1, 3, ten); _o(ws, i + 1, 4, "giờ")
        _o(ws, i + 1, 5, None)
        _o(ws, i + 1, 6, (ds[0].don_gia_dv.get(code) or None) if len(ds) == 1 else None, SO)
    ws.cell(len(LOAI_GIO) + 3, 1, "Đơn giá lương: điền theo hợp đồng (tool không đọc bảng lương của file nguồn)."
            + (" Đơn giá dịch vụ từng hợp đồng: xem cột V trở đi." if len(ds) > 1 else "")).font = F_GHI
    _tieu_de(ws, 1, ["Mã Nhân viên", "Họ tên", "Tỉnh thành", "CMT/CCCD", "Số TK"], bat_dau=9)
    _tieu_de(ws, 1, ["Bộ phận", "Chức danh", "NCC", "Mã thẻ", "Ngày vào", "Ngày nghỉ", "Hợp đồng / Dự án"],
             nen=NEN_PHU, bat_dau=14)
    r = 2
    for da in ds:
        for nv in da.nv.values():
            vals = [nv.ma_nv, nv.ten, nv.tinh, nv.cccd, nv.stk, nv.bo_phan, nv.chuc_danh, nv.ncc, nv.ma_the,
                    nv.ngay_vao, nv.ngay_nghi, da.ma]
            for j, v in enumerate(vals, start=9):
                _o(ws, r, j, v or None, "dd/mm/yyyy" if j in (18, 19) else ("@" if j in (9, 12, 13, 17) else None))
            r += 1
    _tieu_de(ws, 1, ["Hợp đồng / Dự án", "Ký hiệu", "Nội dung", "Đơn giá dịch vụ", "Nguồn"], bat_dau=22)
    r = 2
    for da in ds:
        for code in LOAI_LAM:
            if da.don_gia_dv.get(code):
                for j, v in enumerate([da.ma, code, TEN_LOAI_TUAN[code], da.don_gia_dv[code], da.nguon_don_gia], start=22):
                    _o(ws, r, j, v, SO if j == 25 else None)
                r += 1
    for col, w in {"A": 5, "B": 8, "C": 24, "D": 9, "E": 13, "F": 14, "I": 13, "J": 26, "K": 12, "L": 14,
                   "M": 16, "N": 22, "O": 14, "P": 10, "Q": 11, "R": 11, "S": 11, "T": 12,
                   "V": 12, "W": 8, "X": 24, "Y": 14, "Z": 40}.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    # ================= 4 tab chấm công: tháng / tuần x nhà máy / SK =================
    # Dự án chấm công theo tháng chỉ vào 2 tab tháng, theo tuần chỉ vào 2 tab tuần.
    # Dự án 1 bảng công: cùng dữ liệu ở tab nhà máy và tab SK.
    nhieu = len(ds) > 1
    for ten_tab, lay_cc, kieu in TAB_CHAM_CONG:
        ws = wb.create_sheet(ten_tab)
        r, stt = 1, 0
        cua_tab = [da for da in ds if da.theo_tuan == (kieu == "tuan")]
        for da in cua_tab:
            if kieu == "thang":
                r, stt = _ghi_thang(ws, r, stt, da, lay_cc(da), _tieu_de_hd(da) if nhieu else None)
            else:
                r = _ghi_tuan(ws, r, da, lay_cc(da), nhieu)
        if not cua_tab:
            ws.cell(1, 1, "Không có dự án chấm công theo " + ("tháng" if kieu == "thang" else "tuần")
                    + " trong lượt đồng bộ này").font = F_GHI
            _tieu_de(ws, 2, ["STT", "Mã theo dõi", "Mã NV", "Tên nhân viên",
                             "Tổng công ca ngày" if kieu == "thang" else "Loại công", "TỔNG CỘNG"])
        (_dinh_dang_thang if kieu == "thang" else _dinh_dang_tuan)(ws)

    # ================= 1 tab cho từng dự án (đối chiếu với file gốc) =================
    # Dự án theo tháng -> bố cục bảng tháng, theo tuần -> bố cục bảng tuần. Dự án 2 bảng công: khối công nhà máy
    # rồi khối công SK, ô công SK khác công nhà máy được tô đỏ.
    if nhieu:
        for da in ds:
            ws = wb.create_sheet(ten_tab_du_an(da, wb.sheetnames))
            khoi = [(f"CÔNG NHÀ MÁY – sheet '{da.nguon_nm}'", da.cham_cong_nm, None),
                    (f"CÔNG SK – sheet '{da.nguon_cc}' (ô đỏ: khác công nhà máy)", da.cham_cong, da.cham_cong_nm)]                 if da.hai_bang else [(f"CÔNG NHÀ MÁY = CÔNG SK (dự án chỉ có 1 bảng công: sheet '{da.nguon_cc}')",
                                      da.cham_cong, None)]
            r = 1
            for ten_khoi, cc, khac in khoi:
                td = f"{da.ma} | Kỳ {da.ky} | {ten_khoi}"
                if da.theo_tuan:
                    r = _ghi_tuan(ws, r, da, cc, True, td, khac)
                else:
                    r, _ = _ghi_thang(ws, r, 0, da, cc, td, khac)
            (_dinh_dang_tuan if da.theo_tuan else _dinh_dang_thang)(ws)

    # ================= cham_cong_chi_tiet (dạng dọc cho phân tích) =================
    ws = wb.create_sheet("cham_cong_chi_tiet")
    _tieu_de(ws, 1, ["Hợp đồng / Dự án", "Kỳ lương", "Mã NV", "Tên nhân viên", "Ngày", "Thứ", "Ký hiệu loại giờ",
                     "Loại giờ", "Số giờ", "Ký hiệu ngày (nghỉ/phép...)", "Bảng công", "Sheet nguồn"])
    r = 2
    for da in ds:
        thu_tu = {m: i for i, m in enumerate(da.nv)}
        bang = [("Nhà máy", da.cham_cong_nm, da.nguon_nm), ("SK", da.cham_cong, da.nguon_cc)] if da.hai_bang             else [("Nhà máy = SK", da.cham_cong, da.nguon_cc)]
        for ten_bang, cc, sheet in bang:
            keys = sorted({k for k in cc if k[0] in thu_tu} | {k for k in da.ky_hieu_ngay if k[0] in thu_tu},
                          key=lambda k: (thu_tu[k[0]], k[1]))
            for k in keys:
                g = cc.get(k, {})
                kh = da.ky_hieu_ngay.get(k)
                items = [(c, g[c]) for c in TEN_LOAI if g.get(c)] or [(None, 0)]
                for code, h in items:
                    vals = [da.ma, da.ky, k[0], da.nv[k[0]].ten, dt.datetime.combine(k[1], dt.time()),
                            THU[k[1].weekday()], code, TEN_LOAI_TUAN.get(code) if code else None, h, kh, ten_bang, sheet]
                    for j, v in enumerate(vals, start=1):
                        c = ws.cell(r, j, v)
                        c.font = F_TH
                        if j == 5:
                            c.number_format = "dd/mm/yyyy"
                    r += 1
    for col, w in {"A": 9, "B": 10, "C": 12, "D": 24, "E": 11, "F": 9, "G": 8, "H": 22, "I": 8, "J": 18,
                   "K": 13, "L": 26}.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:L{max(r - 1, 1)}"

    # ================= kiem_tra =================
    ws = wb.create_sheet("kiem_tra")
    td = ["Dự án", "File nguồn", "Biểu mẫu", "Kỳ", "Từ ngày", "Đến ngày", "Chấm công theo", "Số NV",
          "Sheet công nhà máy", "Sheet công SK", "Giờ công nhà máy", "Giờ công SK", "Chênh lệch nhà máy − SK",
          "Số NV lệch nhà máy / SK", "Bảng kê đối chiếu", "Tổng giờ bảng kê", "Chênh lệch nhà máy − bảng kê",
          "Số NV lệch bảng kê", "Số lỗi", "Số nghi ngờ", "Số lưu ý"]
    _tieu_de(ws, 1, td)
    r = 2
    for da in ds:
        ns, dc, dm = doi_chieu_nm_sk(da), doi_chieu(da), dem_muc(da)
        gio_sk = sum(sum(g.get(c, 0) for c in LOAI_LAM) for g in da.tong_gio_nv().values())
        gio_nm = ns["gio_nm"] if ns else gio_sk
        vals = [da.ma, os.path.basename(da.file), da.mau, da.ky, da.tu_ngay, da.den_ngay,
                "Tuần" if da.theo_tuan else "Tháng", len(da.nv),
                da.nguon_nm if da.hai_bang else f"{da.nguon_cc} (dùng chung)", da.nguon_cc, gio_nm, gio_sk,
                gio_nm - gio_sk, (len(ns["nv_lech"]) + len(ns["chi_nm"]) + len(ns["chi_sk"])) if ns else None,
                dc["nguon"] if dc else "không có bảng kê", dc["gio_bk"] if dc else None,
                (gio_nm - dc["gio_bk"]) if dc else None,
                (len(dc["nv_lech"]) + len(dc["chi_cong"]) + len(dc["chi_bk"])) if dc else None,
                dm[LOI], dm[NGHI_NGO], dm[LUU_Y]]
        for j, v in enumerate(vals, start=1):
            if isinstance(v, float):
                v = round(v, 2)
            c = _o(ws, r, j, v, "dd/mm/yyyy" if j in (5, 6) else (GIO if 11 <= j <= 18 else None))
            if j >= 19 and v:
                c.fill = NEN_MUC[[LOI, NGHI_NGO, LUU_Y][j - 19]]
        if ns and not ns["khop"]:
            for j in (13, 14):
                ws.cell(r, j).fill = NEN_DO
        if dc and not dc["khop"]:
            for j in (17, 18):
                ws.cell(r, j).fill = NEN_DO
        r += 1
    r += 2
    if file_loi:
        ws.cell(r, 1, "FILE BỊ BỎ QUA (không đọc được)").font = Font(name=FONT, bold=True, size=11, color="C00000")
        r += 1
        for ten, l in file_loi:
            c = _o(ws, r, 1, "Lỗi"); c.fill = NEN_MUC[LOI]
            _o(ws, r, 2, ten); _o(ws, r, 6, l)
            r += 1
        r += 1
    ws.cell(r, 1, "DANH SÁCH CẢNH BÁO").font = Font(name=FONT, bold=True, size=11)
    r += 1
    for muc, giai_thich in [(LOI, "sai chắc chắn hoặc số liệu không tin được – phải sửa file gốc trước khi dùng"),
                            (NGHI_NGO, "bất thường, cần người kiểm tra lại"),
                            (LUU_Y, "thông tin để biết, không nhất thiết là sai")]:
        c = ws.cell(r, 1, muc)
        c.font, c.fill = F_TD, NEN_MUC[muc]
        ws.cell(r, 2, giai_thich).font = F_GHI
        r += 1
    r += 1
    r_td = r
    _tieu_de(ws, r, ["Mức", "Hợp đồng", "Mã NV", "Tên NV", "Loại", "Nội dung"])
    r += 1
    tat_ca = sorted(((x, da) for da in ds for x in da.van_de), key=lambda t: (THU_TU_MUC[t[0][0]], t[1].ma, t[0][2]))
    for (muc, m, loai, nd), da in tat_ca:
        ten = da.nv[m].ten if m in da.nv else None
        for j, v in enumerate([muc, da.ma, m or None, ten, loai, nd], start=1):
            c = _o(ws, r, j, v)
            if j == 1:
                c.fill = NEN_MUC[muc]
            if j == 6:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    if not tat_ca:
        ws.cell(r, 1, "Không có cảnh báo nào").font = F_TH
    ws.auto_filter.ref = f"A{r_td}:F{max(r - 1, r_td)}"
    for j, w in enumerate([10, 40, 40, 24, 30, 90] + [11] * 12 + [9, 11, 9], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.column_dimensions["D"].width = 24
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = "B2"

    wb.save(out_path)
    log(f"Đã ghi: {out_path}")


def chay(files, out_path, log=print, tach_rieng=False, file_loi=None):
    """file_loi: list (tùy chọn) nhận các (tên file, lỗi) bị bỏ qua."""
    ds = []
    file_loi = [] if file_loi is None else file_loi
    for f in files:
        log(f"Đọc: {os.path.basename(f)}")
        try:
            da = doc_file(f, log)
        except Exception as e:
            file_loi.append((os.path.basename(f), str(e)))
            log(f"  !! BỎ QUA file này: {e}")
            continue
        dm = dem_muc(da)
        dc, ns = doi_chieu(da), doi_chieu_nm_sk(da)
        gio = sum(sum(g.get(c, 0) for c in LOAI_LAM) for g in da.tong_gio_nv().values())
        log(f"  {da.ma} | {da.ky} | {len(da.nv)} NV | "
            + (f"nhà máy '{da.nguon_nm}' {ns['gio_nm']:,.1f} giờ, SK '{da.nguon_cc}' {gio:,.1f} giờ"
               if ns else f"1 bảng công '{da.nguon_cc}' {gio:,.1f} giờ")
            + (f" (bảng kê {dc['gio_bk']:,.1f})" if dc else " (không có bảng kê)") + " | "
            f"{dm[LOI]} lỗi, {dm[NGHI_NGO]} nghi ngờ, {dm[LUU_Y]} lưu ý")
        ds.append(da)
    if not ds:
        raise ValueError("Không đọc được file nào: " + "; ".join(f"{t}: {l}" for t, l in file_loi))
    # 2 file cùng mã hợp đồng -> đổi tên để không ghi đè / gộp nhầm
    dem_ma = Counter(da.ma for da in ds)
    for da in ds:
        if dem_ma[da.ma] > 1:
            cu = da.ma
            da.ma = f"{cu} ({os.path.splitext(os.path.basename(da.file))[0][:30]})"
            da.bao(NGHI_NGO, "", "Trùng mã hợp đồng", f"Nhiều file cùng mã '{cu}' – đã đổi thành '{da.ma}'. "
                   "Kiểm tra có tải trùng file, hoặc đặt tên file chứa 'CTY_<mã>'.")
    ds.sort(key=lambda d: d.ma)
    kiem_tra_cheo(ds)
    tong = {k: sum(dem_muc(da)[k] for da in ds) for k in (LOI, NGHI_NGO, LUU_Y)}
    log(f"Kiểm tra: {tong[LOI]} lỗi, {tong[NGHI_NGO]} nghi ngờ, {tong[LUU_Y]} lưu ý (chi tiết ở sheet kiem_tra)")
    if file_loi:
        log(f"Bỏ qua {len(file_loi)} file lỗi: " + ", ".join(t for t, _ in file_loi))
    ghi_file_chuan(ds, out_path, log, file_loi)
    if tach_rieng:
        goc, _ = os.path.splitext(out_path)
        for da in ds:
            ten = re.sub(r'[\\/:*?"<>|() ]+', "_", da.ma).strip("_")
            ghi_file_chuan([da], f"{goc}_{ten}_{da.ky.replace('/', '-')}.xlsx", log)
    return ds


def main(argv=None):
    ap = argparse.ArgumentParser(description="Đồng bộ bảng chấm công nhiều dự án về 1 file chuẩn "
                                             "(tab tháng / tuần cho công nhà máy và công SK)")
    ap.add_argument("files", nargs="*", help="Các file .xlsx nguồn")
    ap.add_argument("--thu-muc", help="Thư mục chứa các file .xlsx nguồn (lấy tất cả)")
    ap.add_argument("-o", "--output", default=None, help="File kết quả (.xlsx)")
    ap.add_argument("--tach-rieng", action="store_true",
                    help="Xuất thêm 1 file chuẩn riêng cho từng dự án (cột ngày là ngày thật như template)")
    a = ap.parse_args(argv)
    files = list(a.files)
    if a.thu_muc:
        files += [f for f in sorted(glob.glob(os.path.join(a.thu_muc, "*.xlsx")))
                  if not os.path.basename(f).startswith(("~$", "KET_QUA", "Tong_hop"))]
    if not files:
        ap.error("Chưa chọn file nguồn")
    out = a.output or os.path.join(os.path.dirname(os.path.abspath(files[0])),
                                   f"Tong_hop_cham_cong_luong_{dt.datetime.now():%Y%m%d_%H%M}.xlsx")
    chay(files, out, tach_rieng=a.tach_rieng)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    main()
